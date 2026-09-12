from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation



def extract_text(file_path: str) -> str:

    extension = Path(file_path).suffix.lower()


    if extension == ".pdf":

        return extract_pdf(file_path)


    if extension == ".pptx":

        return extract_pptx(file_path)


    raise ValueError(
        "Formato não suportado. Use PDF ou PPTX."
    )





def extract_pdf(file_path: str) -> str:

    reader = PdfReader(file_path)

    text = []


    for page in reader.pages:

        content = page.extract_text()

        if content:
            text.append(content)


    return "\n".join(text)





def extract_pptx(file_path: str) -> str:

    presentation = Presentation(file_path)

    text = []


    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                if shape.text.strip():

                    text.append(shape.text)


    return "\n".join(text)